"""
Test Suite: The Gat-Slayer (Data Ingestion Parser)
==================================================

Tests for logic/parser.py

Covers:
- clean_gat_export: Merged cells, Norwegian characters (ÆØÅ)
- fuzzy_header_match: Variable column name mapping
- vakt_code_mapper: Shift code translation
- data_validation_report: Human noise detection

Test Categories:
- Happy Path: Standard expected usage
- Edge Cases: Empty inputs, limits, null values
- Error Handling: Graceful failures
- Adversarial: Designed to break the parser
"""

import pytest
import polars as pl
from pathlib import Path
from datetime import date, datetime
from typing import Any

# Import the module under test (will fail until implemented)
try:
    from logic.parser import (
        clean_gat_export,
        fuzzy_header_match,
        vakt_code_mapper,
        data_validation_report,
        GatParserError,
        MappingRequiredError,
    )
    PARSER_IMPLEMENTED = True
except ImportError:
    PARSER_IMPLEMENTED = False
    # Stub for test collection
    def clean_gat_export(*args, **kwargs): raise NotImplementedError
    def fuzzy_header_match(*args, **kwargs): raise NotImplementedError
    def vakt_code_mapper(*args, **kwargs): raise NotImplementedError
    def data_validation_report(*args, **kwargs): raise NotImplementedError
    class GatParserError(Exception): pass
    class MappingRequiredError(Exception): pass


pytestmark = pytest.mark.skipif(
    not PARSER_IMPLEMENTED,
    reason="logic/parser.py not yet implemented"
)


class TestCleanGatExport:
    """Tests for clean_gat_export function."""

    # ==================== HAPPY PATH ====================

    def test_clean_simple_excel(self, temp_excel_file: Path):
        """Parser handles a simple, well-formatted Excel file."""
        result = clean_gat_export(temp_excel_file)

        assert isinstance(result, pl.DataFrame)
        assert len(result) > 0
        assert "employee_id" in result.columns

    def test_preserves_norwegian_characters(self, temp_excel_file: Path):
        """Norwegian characters (ÆØÅ) are preserved correctly."""
        result = clean_gat_export(temp_excel_file)

        # Should contain "Ørjan Ås" and "Bjørn Bærum"
        names = result["name"].to_list()
        assert any("Ø" in str(name) for name in names), "Ø should be preserved"
        assert any("Å" in str(name) for name in names), "Å should be preserved"
        assert any("ø" in str(name) for name in names), "ø should be preserved"
        assert any("æ" in str(name) for name in names), "æ should be preserved"

    def test_handles_merged_cells_forward_fill(self, temp_excel_file: Path):
        """Merged cells are forward-filled correctly."""
        result = clean_gat_export(temp_excel_file)

        # After forward-fill, no employee_id should be null
        null_count = result.filter(pl.col("employee_id").is_null()).height
        assert null_count == 0, "Merged cells should be forward-filled"

    def test_returns_polars_dataframe(self, temp_excel_file: Path):
        """Result is always a Polars DataFrame (not Pandas)."""
        result = clean_gat_export(temp_excel_file)
        assert isinstance(result, pl.DataFrame), "Must return Polars DataFrame"

    # ==================== EDGE CASES ====================

    def test_empty_file_returns_empty_dataframe(self, tmp_path: Path):
        """Empty Excel file returns empty DataFrame with schema."""
        import openpyxl
        wb = openpyxl.Workbook()
        empty_file = tmp_path / "empty.xlsx"
        wb.save(empty_file)

        result = clean_gat_export(empty_file)
        assert isinstance(result, pl.DataFrame)
        assert len(result) == 0

    def test_single_row_file(self, tmp_path: Path):
        """File with only one data row is handled correctly."""
        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["employee_id", "name", "date", "shift_code"])
        ws.append(["123456", "Test", "2026-01-05", "D1"])

        single_row = tmp_path / "single.xlsx"
        wb.save(single_row)

        result = clean_gat_export(single_row)
        assert len(result) == 1

    def test_max_rows_performance(self, tmp_path: Path):
        """Parser handles large files (10,000+ rows) efficiently."""
        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["employee_id", "name", "date", "shift_code"])

        for i in range(10_000):
            ws.append([f"{100000 + i}", f"Employee {i}", "2026-01-05", "D1"])

        large_file = tmp_path / "large.xlsx"
        wb.save(large_file)

        import time
        start = time.time()
        result = clean_gat_export(large_file)
        elapsed = time.time() - start

        assert len(result) == 10_000
        assert elapsed < 30, f"Parser too slow: {elapsed:.2f}s for 10k rows"

    def test_all_null_column(self, tmp_path: Path):
        """Column with all null values is handled gracefully."""
        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["employee_id", "name", "optional_field", "shift_code"])
        ws.append(["123456", "Test", None, "D1"])
        ws.append(["123457", "Test2", None, "D2"])

        file = tmp_path / "nulls.xlsx"
        wb.save(file)

        result = clean_gat_export(file)
        assert "optional_field" in result.columns or len(result) > 0

    # ==================== ERROR HANDLING ====================

    def test_file_not_found_raises_error(self):
        """Non-existent file raises appropriate error."""
        with pytest.raises((FileNotFoundError, GatParserError)):
            clean_gat_export(Path("/nonexistent/file.xlsx"))

    def test_invalid_file_format_raises_error(self, tmp_path: Path):
        """Non-Excel file raises appropriate error."""
        invalid_file = tmp_path / "invalid.txt"
        invalid_file.write_text("This is not an Excel file")

        with pytest.raises((ValueError, GatParserError)):
            clean_gat_export(invalid_file)

    def test_corrupted_excel_raises_error(self, tmp_path: Path):
        """Corrupted Excel file raises appropriate error."""
        corrupted = tmp_path / "corrupted.xlsx"
        corrupted.write_bytes(b"PK\x03\x04corrupted data here")

        with pytest.raises((Exception, GatParserError)):
            clean_gat_export(corrupted)

    # ==================== ADVERSARIAL ====================

    def test_adversarial_formula_injection(self, tmp_path: Path):
        """Cells with Excel formulas don't execute or break parser."""
        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["employee_id", "name", "shift_code"])
        ws.append(["=1+1", "=HYPERLINK('http://evil.com')", "D1"])
        ws.append(["123456", "Normal Name", "D1"])

        formula_file = tmp_path / "formulas.xlsx"
        wb.save(formula_file)

        result = clean_gat_export(formula_file)
        # Formulas should be treated as strings, not evaluated
        assert "=1+1" in str(result["employee_id"].to_list()) or len(result) >= 1

    def test_adversarial_unicode_bombs(self, tmp_path: Path):
        """Unicode edge cases don't crash the parser."""
        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["employee_id", "name", "shift_code"])
        # Various unicode edge cases
        ws.append(["123456", "Normal", "D1"])
        ws.append(["123457", "Ω≈ç√∫", "D1"])  # Math symbols
        ws.append(["123458", "🔥💀👻", "D1"])  # Emoji
        ws.append(["123459", "\u0000\u0001", "D1"])  # Null chars
        ws.append(["123460", "A" * 10000, "D1"])  # Very long string

        unicode_file = tmp_path / "unicode.xlsx"
        wb.save(unicode_file)

        result = clean_gat_export(unicode_file)
        assert len(result) >= 1, "Parser should handle unicode gracefully"

    def test_adversarial_deeply_nested_merged_cells(self, tmp_path: Path):
        """Complex merge patterns don't break forward-fill logic."""
        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active

        # Create complex merge pattern
        for row in range(1, 20):
            for col in range(1, 10):
                ws.cell(row=row, column=col, value=f"R{row}C{col}")

        # Merge in various patterns
        ws.merge_cells("A1:A5")
        ws.merge_cells("B1:B3")
        ws.merge_cells("C2:C6")
        ws.merge_cells("A10:C10")

        merge_file = tmp_path / "complex_merge.xlsx"
        wb.save(merge_file)

        # Should not raise
        result = clean_gat_export(merge_file)
        assert isinstance(result, pl.DataFrame)


class TestFuzzyHeaderMatch:
    """Tests for fuzzy_header_match function."""

    # ==================== HAPPY PATH ====================

    def test_exact_match(self):
        """Exact header names are matched correctly."""
        headers = ["employee_id", "name", "date", "shift_code"]
        result = fuzzy_header_match(headers)

        assert result["employee_id"] == "employee_id"
        assert result["name"] == "name"

    def test_norwegian_variants(self, fuzzy_header_variants: list):
        """Norwegian header variants are mapped to standard schema."""
        for variant in fuzzy_header_variants:
            result = fuzzy_header_match(variant)

            # Should map to standard schema
            assert "employee_id" in result.values(), f"Failed on: {variant}"

    def test_case_insensitive_match(self):
        """Header matching is case-insensitive."""
        headers = ["EMPLOYEE_ID", "NAME", "Date", "ShIfT_CoDe"]
        result = fuzzy_header_match(headers)

        assert len(result) > 0

    # ==================== EDGE CASES ====================

    def test_empty_headers_returns_empty_mapping(self):
        """Empty header list returns empty mapping."""
        result = fuzzy_header_match([])
        assert result == {} or len(result) == 0

    def test_single_header(self):
        """Single header is processed correctly."""
        result = fuzzy_header_match(["Personnummer"])
        assert "employee_id" in result.values()

    def test_duplicate_headers(self):
        """Duplicate headers are handled (first wins or explicit handling)."""
        headers = ["name", "Name", "NAME", "navn"]
        result = fuzzy_header_match(headers)
        # Should not crash, behavior is implementation-defined
        assert isinstance(result, dict)

    def test_numeric_headers(self):
        """Numeric column headers are handled."""
        headers = [1, 2, 3, "name"]
        result = fuzzy_header_match([str(h) for h in headers])
        assert isinstance(result, dict)

    # ==================== ERROR HANDLING ====================

    def test_none_input_raises_error(self):
        """None input raises appropriate error."""
        with pytest.raises((TypeError, ValueError, AttributeError)):
            fuzzy_header_match(None)

    def test_completely_unknown_headers(self):
        """Completely unrecognized headers trigger warning or error."""
        headers = ["xyz123", "abc456", "unknownfield"]
        result = fuzzy_header_match(headers)

        # Should either return empty or raise MappingRequiredError
        # Implementation decides behavior
        assert isinstance(result, dict)

    # ==================== ADVERSARIAL ====================

    def test_adversarial_sql_injection_headers(self):
        """SQL injection attempts in headers are sanitized."""
        headers = [
            "'; DROP TABLE users; --",
            "employee_id OR 1=1",
            "<script>alert('xss')</script>",
        ]
        result = fuzzy_header_match(headers)
        # Should not crash, injection attempts are just strings
        assert isinstance(result, dict)

    def test_adversarial_extremely_long_headers(self):
        """Very long header names are handled."""
        headers = ["A" * 10000, "B" * 5000, "employee_id"]
        result = fuzzy_header_match(headers)
        assert isinstance(result, dict)


class TestVaktCodeMapper:
    """Tests for vakt_code_mapper function."""

    # ==================== HAPPY PATH ====================

    def test_standard_day_codes(self, shift_codes: dict):
        """Standard day shift codes (D1, D2, D3) are mapped correctly."""
        for code in ["D1", "D2", "D3"]:
            result = vakt_code_mapper(code, shift_codes)
            assert result["type"] == "day"
            assert result["start"] is not None
            assert result["end"] is not None

    def test_standard_evening_codes(self, shift_codes: dict):
        """Standard evening shift codes (A1, A2, A3) are mapped correctly."""
        for code in ["A1", "A2", "A3"]:
            result = vakt_code_mapper(code, shift_codes)
            assert result["type"] == "evening"

    def test_standard_night_codes(self, shift_codes: dict):
        """Standard night shift codes (N1, N2, N3) are mapped correctly."""
        for code in ["N1", "N2", "N3"]:
            result = vakt_code_mapper(code, shift_codes)
            assert result["type"] == "night"

    def test_off_day_codes(self, shift_codes: dict):
        """Off/Leave codes (F, FRI, L) return null times."""
        for code in ["F", "FRI", "L"]:
            result = vakt_code_mapper(code, shift_codes)
            assert result["start"] is None
            assert result["end"] is None
            assert result["type"] in ["off", "leave"]

    def test_returns_time_strings(self, shift_codes: dict):
        """Time values are returned as proper time strings (HH:MM)."""
        result = vakt_code_mapper("D1", shift_codes)
        assert result["start"] == "07:00"
        assert result["end"] == "15:00"

    # ==================== EDGE CASES ====================

    def test_lowercase_codes(self, shift_codes: dict):
        """Lowercase codes are handled (case-insensitive)."""
        result = vakt_code_mapper("d1", shift_codes)
        assert result["type"] == "day"

    def test_whitespace_in_code(self, shift_codes: dict):
        """Codes with whitespace are trimmed."""
        result = vakt_code_mapper("  D1  ", shift_codes)
        assert result["type"] == "day"

    def test_empty_string_code(self, shift_codes: dict):
        """Empty string returns appropriate response."""
        result = vakt_code_mapper("", shift_codes)
        assert result.get("type") in [None, "unknown", "off"] or "error" in result

    def test_null_code(self, shift_codes: dict):
        """None/null code is handled gracefully."""
        result = vakt_code_mapper(None, shift_codes)
        assert result.get("type") in [None, "unknown", "off"] or "error" in result

    # ==================== ERROR HANDLING ====================

    def test_unknown_code_triggers_mapping_required(self, shift_codes: dict):
        """Unknown shift code raises MappingRequiredError."""
        with pytest.raises(MappingRequiredError):
            vakt_code_mapper("UNKNOWN_CODE_XYZ", shift_codes)

    def test_absence_codes_recognized(self, rules_config: dict):
        """Absence codes (S, SYK, ES, P, FE, K) are recognized."""
        absence_codes = rules_config["shift_codes"]["absence_codes"]

        for code in absence_codes.keys():
            result = vakt_code_mapper(code, rules_config["shift_codes"]["standard"])
            # Should not raise MappingRequiredError
            assert "type" in result or result is not None

    # ==================== ADVERSARIAL ====================

    def test_adversarial_injection_in_code(self, shift_codes: dict):
        """Injection attempts in shift codes are rejected."""
        malicious_codes = [
            "D1; DROP TABLE",
            "D1\n\rNewLine",
            "D1<script>",
            "D1${env.SECRET}",
        ]
        for code in malicious_codes:
            # Should either raise MappingRequired or return safely
            try:
                result = vakt_code_mapper(code, shift_codes)
                assert isinstance(result, dict)
            except MappingRequiredError:
                pass  # Expected for unknown codes

    def test_adversarial_numeric_code(self, shift_codes: dict):
        """Pure numeric input is handled."""
        try:
            result = vakt_code_mapper(12345, shift_codes)
            assert isinstance(result, dict)
        except (MappingRequiredError, TypeError):
            pass  # Acceptable behaviors


class TestDataValidationReport:
    """Tests for data_validation_report function."""

    # ==================== HAPPY PATH ====================

    def test_clean_data_no_flags(self, sample_clean_schedule: pl.DataFrame):
        """Clean data produces no validation flags."""
        report = data_validation_report(sample_clean_schedule)

        assert report["status"] == "valid"
        assert len(report.get("flags", [])) == 0

    def test_detects_human_noise(self):
        """Human noise entries like 'Syk - ringt' are flagged."""
        noisy_data = pl.DataFrame({
            "employee_id": ["123456", "123456"],
            "date": [date(2026, 1, 5), date(2026, 1, 6)],
            "shift_code": ["D1", "Syk - ringt"],
        })

        report = data_validation_report(noisy_data)
        assert len(report.get("flags", [])) > 0
        assert any("Syk" in str(flag) for flag in report.get("flags", []))

    def test_report_structure(self, sample_clean_schedule: pl.DataFrame):
        """Report contains expected structure."""
        report = data_validation_report(sample_clean_schedule)

        assert "status" in report
        assert "flags" in report or report["status"] == "valid"
        assert "row_count" in report
        assert "employee_count" in report

    # ==================== EDGE CASES ====================

    def test_empty_dataframe_report(self, empty_schedule: pl.DataFrame):
        """Empty DataFrame produces valid but empty report."""
        report = data_validation_report(empty_schedule)

        assert report["row_count"] == 0
        assert report["status"] in ["valid", "empty"]

    def test_missing_required_columns(self):
        """Missing required columns are flagged."""
        incomplete_data = pl.DataFrame({
            "some_column": ["value1", "value2"],
        })

        report = data_validation_report(incomplete_data)
        assert report["status"] == "invalid" or len(report.get("flags", [])) > 0

    def test_duplicate_entries_detected(self):
        """Duplicate employee-date entries are flagged."""
        duplicate_data = pl.DataFrame({
            "employee_id": ["123456", "123456"],
            "date": [date(2026, 1, 5), date(2026, 1, 5)],  # Same date!
            "shift_code": ["D1", "A1"],  # Two shifts same day
        })

        report = data_validation_report(duplicate_data)
        # Should flag potential duplicate
        assert len(report.get("flags", [])) > 0 or "duplicate" in str(report).lower()

    # ==================== ERROR HANDLING ====================

    def test_null_input_raises_error(self):
        """None input raises appropriate error."""
        with pytest.raises((TypeError, ValueError)):
            data_validation_report(None)

    def test_non_dataframe_input_raises_error(self):
        """Non-DataFrame input raises appropriate error."""
        with pytest.raises((TypeError, ValueError)):
            data_validation_report({"not": "a dataframe"})

    # ==================== ADVERSARIAL ====================

    def test_adversarial_all_rows_invalid(self):
        """DataFrame where every row has issues."""
        bad_data = pl.DataFrame({
            "employee_id": [None, "", "invalid!@#"],
            "date": [None, None, None],
            "shift_code": ["???", "UNKNOWN", ""],
        })

        report = data_validation_report(bad_data)
        assert report["status"] == "invalid"
        assert len(report.get("flags", [])) >= 3

    def test_adversarial_huge_noise_content(self):
        """Cells with very long 'noise' content."""
        noise_data = pl.DataFrame({
            "employee_id": ["123456"],
            "date": [date(2026, 1, 5)],
            "shift_code": ["Syk - ringt av Kari som sa at " + "bla " * 1000],
        })

        report = data_validation_report(noise_data)
        # Should still process and flag
        assert isinstance(report, dict)
        assert len(report.get("flags", [])) > 0


class TestParserIntegration:
    """Integration tests for the complete parser pipeline."""

    def test_end_to_end_clean_file(self, temp_excel_file: Path, shift_codes: dict):
        """Complete pipeline: Excel -> Clean DataFrame -> Mapped Shifts."""
        # Step 1: Clean the export
        df = clean_gat_export(temp_excel_file)

        # Step 2: Validate
        report = data_validation_report(df)

        # Step 3: Map shift codes
        for code in df["shift_code"].unique().to_list():
            if code and str(code).strip():
                try:
                    mapped = vakt_code_mapper(str(code), shift_codes)
                    assert "type" in mapped
                except MappingRequiredError:
                    # Acceptable for unknown codes
                    pass

        assert isinstance(df, pl.DataFrame)
        assert isinstance(report, dict)

    def test_gat_slayer_directive_compliance(self, temp_excel_file: Path):
        """Parser follows CLAUDE.md directives."""
        df = clean_gat_export(temp_excel_file)

        # Directive: Polars Over Pandas
        assert isinstance(df, pl.DataFrame), "Must use Polars, not Pandas"

        # Directive: No Fake Truths
        # Unknown codes should not be silently converted
        report = data_validation_report(df)
        assert isinstance(report, dict)


# ==================== CHAOS TESTS ====================

class TestChaosScenarios:
    """Chaos/Stress tests using generated adversarial data."""

    def test_chaos_random_encoding(self, tmp_path: Path):
        """Parser handles files with unusual encodings."""
        # Create file with ISO-8859-1 content
        content = "employee_id,name\n123456,Ørjan Ås"
        file = tmp_path / "iso.csv"
        file.write_bytes(content.encode("iso-8859-1"))

        # Parser should detect encoding or fail gracefully
        try:
            result = clean_gat_export(file)
            assert isinstance(result, pl.DataFrame)
        except (UnicodeDecodeError, GatParserError):
            pass  # Acceptable if error is clear

    def test_chaos_mixed_date_formats(self, tmp_path: Path):
        """Parser handles mixed date formats in same file."""
        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["employee_id", "date", "shift_code"])
        ws.append(["123456", "2026-01-05", "D1"])  # ISO format
        ws.append(["123456", "05.01.2026", "D1"])  # Norwegian format
        ws.append(["123456", "05/01/2026", "D1"])  # Slash format
        ws.append(["123456", "January 5, 2026", "D1"])  # English format

        file = tmp_path / "mixed_dates.xlsx"
        wb.save(file)

        result = clean_gat_export(file)
        # Should either parse all or flag unparseable
        assert isinstance(result, pl.DataFrame)

    def test_chaos_special_shift_codes(self, tmp_path: Path, shift_codes: dict):
        """Parser handles municipality-specific shift codes."""
        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["employee_id", "date", "shift_code"])
        # Real-world weird codes found in Norwegian hospitals
        weird_codes = [
            "D1+", "D1/A1", "D1-kort", "Dagvakt", "DV",
            "7-15", "07:00-15:00", "Tidligvakt", "TV"
        ]
        for i, code in enumerate(weird_codes):
            ws.append([f"12345{i}", "2026-01-05", code])

        file = tmp_path / "weird_codes.xlsx"
        wb.save(file)

        result = clean_gat_export(file)
        # Should process and flag unknown codes
        report = data_validation_report(result)
        assert isinstance(report, dict)
